from pathlib import Path
from uuid import uuid4

from docx import Document
from fastapi import UploadFile
from pypdf import PdfReader

from app.config import get_settings


def save_upload(upload: UploadFile, subdir: str = "resumes") -> tuple[str, Path]:
    settings = get_settings()
    extension = Path(upload.filename or "").suffix.lower()
    filename = f"{uuid4().hex}{extension}"
    target = settings.upload_dir / subdir / filename
    content = upload.file.read()
    target.write_bytes(content)
    upload.file.seek(0)
    return filename, target


def extract_text(file_path: Path) -> str:
    suffix = file_path.suffix.lower()
    if suffix == ".pdf":
        reader = PdfReader(str(file_path))
        text = "\n".join(page.extract_text() or "" for page in reader.pages)
        return text.strip()
    if suffix == ".docx":
        doc = Document(str(file_path))
        return "\n".join(paragraph.text for paragraph in doc.paragraphs if paragraph.text.strip()).strip()
    if suffix in {".txt", ".md"}:
        return file_path.read_text(encoding="utf-8")
    raise ValueError("Unsupported resume format. Please upload PDF, DOCX, TXT, or MD.")


def build_docx_resume(output_path: Path, title: str, summary: str, body_text: str) -> None:
    doc = Document()
    doc.add_heading(title, level=0)
    doc.add_paragraph(summary)
    for block in [part.strip() for part in body_text.split("\n\n") if part.strip()]:
        doc.add_paragraph(block)
    doc.save(output_path)

